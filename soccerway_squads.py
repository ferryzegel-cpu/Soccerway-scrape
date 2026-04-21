"""
Soccerway Dutch Leagues Squad Scraper
======================================
Scrapt spelersstatistieken van alle clubs in de Eredivisie en
Keuken Kampioen Divisie en slaat ze op als Excel-bestand (één tabblad per club).

Club-URL's worden automatisch opgehaald van de competitiestandenpagina's.

Gebruik:
    python soccerway_squads.py
    python soccerway_squads.py --output mijn_bestand.xlsx
    python soccerway_squads.py --competities eredivisie kkd

Installeer benodigdheden:
    pip install requests beautifulsoup4 lxml openpyxl
"""

import argparse
import re
import sys
import time

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Configuratie
# ---------------------------------------------------------------------------

BASE_URL = "https://www.soccerway.com"

COMPETITIES = {
    "eredivisie": {
        "naam": "Eredivisie",
        "standings_url": f"{BASE_URL}/netherlands/eredivisie/standings/",
        "kleur": "1E3A5F",
        "accent": "FF6B00",
    },
    "kkd": {
        "naam": "Keuken Kampioen Divisie",
        "standings_url": f"{BASE_URL}/netherlands/eerste-divisie/standings/",
        "kleur": "2D5016",
        "accent": "FFD700",
    },
}

KOLOMMEN = ["#", "Naam", "Positie", "Leeftijd", "Wedstrijden", "MIN",
            "Goals", "Assists", "Gele kaarten", "Rode kaarten"]

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "nl-NL,nl;q=0.9,en;q=0.8",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Referer": BASE_URL + "/",
}

REQUEST_DELAY = 2

POSITIE_MAP = {
    "Goalkeepers": "Keeper",
    "Defenders":   "Verdediger",
    "Midfielders": "Middenvelder",
    "Forwards":    "Aanvaller",
    "Coach":       "Coach",
}


# ---------------------------------------------------------------------------
# Hulpfuncties
# ---------------------------------------------------------------------------

def fetch(url: str, session: requests.Session):
    try:
        r = session.get(url, headers=HEADERS, timeout=20)
        r.raise_for_status()
        return BeautifulSoup(r.text, "lxml")
    except requests.RequestException as e:
        print(f"  [WAARSCHUWING] {url}: {e}", file=sys.stderr)
        return None


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip()


# ---------------------------------------------------------------------------
# Club-URL's automatisch ophalen
# ---------------------------------------------------------------------------

def haal_clubs_op(comp_key: str, session: requests.Session) -> list:
    """
    Haalt club-URLs op via de soccerway standenspagina.
    Zoekt naar alle /team/slug/hash/ links op de pagina.
    """
    cfg = COMPETITIES[comp_key]
    url = cfg["standings_url"]
    print(f"  Clublijst ophalen: {url}")

    soup = fetch(url, session)
    if not soup:
        return []

    clubs = []
    gezien = set()

    for a in soup.find_all("a", href=True):
        href = a["href"]
        # Soccerway team-links: /team/naam/HASH/ (hash = 8 alfanumerieke tekens)
        m = re.search(r"/team/([^/]+)/([A-Za-z0-9]{6,10})/?$", href)
        if m:
            slug, hsh = m.group(1), m.group(2)
            squad_url = f"{BASE_URL}/team/{slug}/{hsh}/squad/"
            naam = clean(a.get_text())
            if squad_url not in gezien and naam and len(naam) > 1 and not naam.isdigit():
                gezien.add(squad_url)
                clubs.append((naam, squad_url))

    # Fallback: probeer de transfers-pagina als standenpagina leeg is
    if not clubs:
        fallback = url.replace("/standings/", "/transfers/")
        print(f"  Geen clubs via standings, probeer: {fallback}", file=sys.stderr)
        soup2 = fetch(fallback, session)
        if soup2:
            for a in soup2.find_all("a", href=True):
                href = a["href"]
                m = re.search(r"/team/([^/]+)/([A-Za-z0-9]{6,10})/?$", href)
                if m:
                    slug, hsh = m.group(1), m.group(2)
                    squad_url = f"{BASE_URL}/team/{slug}/{hsh}/squad/"
                    naam = clean(a.get_text())
                    if squad_url not in gezien and naam and len(naam) > 1 and not naam.isdigit():
                        gezien.add(squad_url)
                        clubs.append((naam, squad_url))

    print(f"  {len(clubs)} club(s) gevonden")
    return clubs


# ---------------------------------------------------------------------------
# Spelersdata scrapen
# ---------------------------------------------------------------------------

def scrape_squad(url: str, session: requests.Session) -> list:
    soup = fetch(url, session)
    if not soup:
        return []

    spelers = []
    current_positie = "Onbekend"

    for row in soup.select("table.statistics tr, table tr"):
        th = row.find("th")
        if th:
            tekst = clean(th.get_text())
            if tekst in POSITIE_MAP:
                current_positie = POSITIE_MAP[tekst]
            continue

        if current_positie == "Coach":
            continue

        tds = row.find_all("td")
        if len(tds) < 3:
            continue

        def td(i):
            return clean(tds[i].get_text()) if i < len(tds) else "-"

        naam_td = tds[1] if len(tds) > 1 else None
        if not naam_td:
            continue
        a_tag = naam_td.find("a")
        naam = clean(a_tag.get_text() if a_tag else naam_td.get_text())

        if not naam or naam in ("#", "Name", "") or naam.isdigit():
            continue

        try:
            spelers.append({
                "#":            td(0),
                "Naam":         naam,
                "Positie":      current_positie,
                "Leeftijd":     td(2),
                "Wedstrijden":  td(3),
                "MIN":          td(4),
                "Goals":        td(5),
                "Assists":      td(6),
                "Gele kaarten": td(7),
                "Rode kaarten": td(8),
            })
        except (IndexError, AttributeError):
            continue

    return spelers


# ---------------------------------------------------------------------------
# Excel opmaak
# ---------------------------------------------------------------------------

def maak_border():
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def schrijf_tabblad(wb, clubnaam, comp_cfg, spelers, tab_kleur):
    tab_naam = re.sub(r"[/\\?*:\[\]]", "", clubnaam)[:31]
    if tab_naam in [ws.title for ws in wb.worksheets]:
        tab_naam = tab_naam[:28] + " (2)"

    ws = wb.create_sheet(title=tab_naam)
    ws.sheet_properties.tabColor = tab_kleur
    border = maak_border()

    ws.merge_cells("A1:J1")
    ws["A1"] = clubnaam
    ws["A1"].font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", start_color=comp_cfg["kleur"])
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:J2")
    ws["A2"] = comp_cfg["naam"]
    ws["A2"].font = Font(name="Arial", italic=True, size=10, color="FFFFFF")
    ws["A2"].fill = PatternFill("solid", start_color=comp_cfg["accent"])
    ws["A2"].alignment = Alignment(horizontal="center")

    for col, kol_naam in enumerate(KOLOMMEN, start=1):
        cel = ws.cell(row=3, column=col, value=kol_naam)
        cel.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cel.fill = PatternFill("solid", start_color="2C3E50")
        cel.alignment = Alignment(horizontal="center")
        cel.border = border

    pos_kleur = {
        "Keeper": "E8F4FD", "Verdediger": "E8F8E8",
        "Middenvelder": "FFF8E1", "Aanvaller": "FDE8E8", "Onbekend": "F5F5F5",
    }
    for rij_idx, speler in enumerate(spelers, start=4):
        vul = PatternFill("solid", start_color=pos_kleur.get(speler.get("Positie", "Onbekend"), "FFFFFF"))
        for col, kol_naam in enumerate(KOLOMMEN, start=1):
            cel = ws.cell(row=rij_idx, column=col, value=speler.get(kol_naam, "-"))
            cel.font = Font(name="Arial", size=10)
            cel.fill = vul
            cel.border = border
            cel.alignment = Alignment(horizontal="center" if col != 2 else "left")

    for i, b in enumerate([6, 30, 16, 10, 14, 8, 8, 9, 14, 13], start=1):
        ws.column_dimensions[get_column_letter(i)].width = b

    ws.freeze_panes = "A4"


def maak_overzicht_tab(wb, club_data):
    ws = wb.active
    ws.title = "Overzicht"
    ws.sheet_properties.tabColor = "1A1A2E"

    ws.merge_cells("A1:D1")
    ws["A1"] = "Nederlandse Voetbal — Spelersdata Overzicht"
    ws["A1"].font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", start_color="1A1A2E")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 28

    for col, h in enumerate(["Club", "Competitie", "Spelers", "Tabblad"], 1):
        cel = ws.cell(row=2, column=col, value=h)
        cel.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cel.fill = PatternFill("solid", start_color="2C3E50")
        cel.alignment = Alignment(horizontal="center")

    for rij, info in enumerate(club_data, start=3):
        ws.cell(rij, 1, info["club"]).font = Font(name="Arial", size=10)
        ws.cell(rij, 2, info["competitie"]).font = Font(name="Arial", size=10)
        cel_n = ws.cell(rij, 3, info["aantal"])
        cel_n.font = Font(name="Arial", size=10)
        cel_n.alignment = Alignment(horizontal="center")
        ws.cell(rij, 4, info["tab"]).font = Font(name="Arial", size=10)
        if rij % 2 == 0:
            for col in range(1, 5):
                ws.cell(rij, col).fill = PatternFill("solid", start_color="F0F4FF")

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 28


# ---------------------------------------------------------------------------
# Hoofdprogramma
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Soccerway squad scraper → Excel")
    parser.add_argument("--output", "-o", default="nederland_voetbal_squads.xlsx")
    parser.add_argument("--competities", "-c", nargs="+",
                        choices=["eredivisie", "kkd"], default=["eredivisie", "kkd"])
    parser.add_argument("--vertraging", "-v", type=float, default=REQUEST_DELAY)
    args = parser.parse_args()

    session = requests.Session()
    wb = Workbook()
    club_data_overzicht = []

    print(f"\nSoccerway Squad Scraper")
    print("=" * 50)
    print(f"Output: {args.output}\n")

    for comp_key in args.competities:
        cfg = COMPETITIES[comp_key]
        print(f"\n▶ {cfg['naam']}")

        clubs = haal_clubs_op(comp_key, session)
        time.sleep(args.vertraging)

        if not clubs:
            print(f"  [OVERGESLAGEN] Geen clubs gevonden")
            continue

        print(f"\n  Squads scrapen ({len(clubs)} clubs):")
        for i, (clubnaam, url) in enumerate(clubs, 1):
            print(f"  [{i}/{len(clubs)}] {clubnaam} ...", end=" ", flush=True)
            spelers = scrape_squad(url, session)

            if spelers:
                schrijf_tabblad(wb, clubnaam, cfg, spelers, cfg["kleur"])
                print(f"✓ {len(spelers)} spelers")
            else:
                print("⚠ Geen data")

            club_data_overzicht.append({
                "club":       clubnaam,
                "competitie": cfg["naam"],
                "aantal":     len(spelers),
                "tab":        re.sub(r"[/\\?*:\[\]]", "", clubnaam)[:31],
            })

            if i < len(clubs):
                time.sleep(args.vertraging)

    maak_overzicht_tab(wb, club_data_overzicht)
    wb.save(args.output)

    totaal = sum(d["aantal"] for d in club_data_overzicht)
    print(f"\n{'=' * 50}")
    print(f"✓ Opgeslagen: {args.output}")
    print(f"  {len(wb.worksheets)} tabbladen  |  {totaal} spelers totaal")


if __name__ == "__main__":
    main()
